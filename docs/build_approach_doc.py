"""Generates the team approach document (docs/Approach_Pitch.docx).

Kept as a script so the document can be regenerated after the numbers change,
rather than being hand edited and drifting out of sync with the pipeline.

    python docs/build_approach_doc.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Approach_Pitch.docx"

ACCENT = RGBColor(0x1F, 0x4F, 0x82)
GREY = RGBColor(0x5C, 0x66, 0x73)


def style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12


def head(doc, text, size=13, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    return p


def para(doc, text, bold=False, italic=False, size=10.5, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def build():
    doc = Document()
    style(doc)

    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Pt(46)
    sec.left_margin = sec.right_margin = Pt(52)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("Business Entity Resolution: our approach")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run("Amazon ML Challenge 2026, Round 1. Written for the team before we split the work.")
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY

    head(doc, "1. What the task actually is", space_before=4)
    para(doc, "We get business records from three sources. The same shop can appear in all "
              "three, spelled differently each time, with no shared ID linking them. Source 1 "
              "is the clean master list. For every Source 1 record, we have to output every "
              "Source 2 and Source 3 record that is the same business. It can be none, one, "
              "or several. The problem statement says the average is closer to several.")

    head(doc, "2. How the score works, and why it changes our strategy")
    para(doc, "The metric is F0.5, averaged separately for each Source 1 record. The formula "
              "is given in the problem statement. The important part is that precision counts "
              "twice as much as recall, so a wrong match hurts us more than a missed one.")
    para(doc, "Worked through on one record, that means:")
    bullet(doc, "we get the full 1.0", bold_prefix="Right match, nothing extra: ")
    bullet(doc, "we drop to about 0.56", bold_prefix="Right match plus one wrong one: ")
    bullet(doc, "we get 0.0", bold_prefix="Correct match missed entirely: ")
    bullet(doc, "we lose the whole point for that record", bold_prefix="Any guess on a record that truly has no match: ")
    para(doc, "So when the model is unsure, staying quiet is usually the better bet. That one "
              "idea drives most of the design below.", bold=True)

    head(doc, "3. Four things we measured in the data")
    para(doc, "These are not assumptions. They come from our own profiling script "
              "(src/profile_data.py), run over the full training set.")
    bullet(doc, "Only 5.58 percent of Source 1 records have no match at all. Submitting nothing "
                "for everyone would score 0.0558, so that is our floor, not a plan.",
           bold_prefix="Singletons are rare. ")
    bullet(doc, "No Source 2 or Source 3 record is ever claimed by two different Source 1 "
                "records. We checked all 7,638,365 links and found zero exceptions. We can "
                "therefore force our predictions to be exclusive, which costs a little recall "
                "and buys precision. Precision is what the metric pays for.",
           bold_prefix="Each record has at most one owner. ")
    bullet(doc, "Across 1,038,755 sampled links, not one match crossed a country label. That "
                "makes country a free filter that cuts the search space by roughly 2.6 times "
                "with no loss at all.",
           bold_prefix="Matches never cross countries. ")
    bullet(doc, "Source 1 is always in Latin script, but about 10 percent of the Indian matches "
                "point at a Devanagari record. Character and word overlap score exactly zero "
                "across two different scripts, so those matches are invisible unless we convert "
                "first. We wrote our own converter, since the rules forbid outside services.",
           bold_prefix="Roughly one in ten Indian matches is in a different script. ")

    head(doc, "4. The pipeline, five steps")
    para(doc, "The shape of this is standard entity resolution, and the problem statement "
              "points at it directly in its own Tips for Success section, which tells us to "
              "invest in blocking and to look at Jaccard, Levenshtein and TF-IDF cosine.")

    steps = [
        ("Clean the text",
         "Lowercase, drop punctuation, strip accents so the French records line up, convert "
         "Devanagari to Latin, and fold abbreviations so Pvt Ltd and Private Limited become "
         "the same thing. Street numbers and PIN codes get pulled into their own field, since "
         "they separate businesses far better than words like Road do."),
        ("Blocking, also called candidate generation",
         "We cannot compare 1.7 million records against 10 million, that is 17 trillion pairs. "
         "So for each record we shortlist a handful of plausible partners using TF-IDF cosine "
         "on the name, within the same country. This step decides the highest recall we can "
         "ever reach, because anything it misses is gone for good. Amazon asks us to submit "
         "this candidate list as a separate file and reviews its quality, so it is graded work, "
         "not scratch work."),
        ("Build comparison features",
         "For each shortlisted pair, measure how alike the two records are: name similarity, "
         "address similarity, how many rare words they share, whether the house number and PIN "
         "code match. Sharing an unusual word like Zephay tells us much more than sharing the "
         "word Restaurant, so rare words get weighted higher."),
        ("Train a classifier",
         "Feed those features to a gradient boosting model, LightGBM. It learns which "
         "combinations mean same business. The negative examples have to be the near misses "
         "our own blocking produced, not random pairs, because random pairs are too easy and "
         "teach it nothing."),
        ("Pick the cut off, then enforce exclusivity",
         "Sweep the accept threshold against F0.5 on our own held out data and take the best "
         "one. Expect it to sit high, because the metric punishes false matches. Then apply "
         "finding number two: if two Source 1 records both claim the same partner, only the "
         "stronger claim survives."),
    ]
    for i, (name, body) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"Step {i}. {name}. ")
        r.bold = True
        p.add_run(body)

    head(doc, "5. Why we are building it in this order")
    para(doc, "Steps 1 and 2 block everything else, so they go first and they cannot slip. "
              "A weak shortlist caps the final score no matter how good the model is. We also "
              "want a simple rules only version submitted on day one, before any machine "
              "learning, because it proves our output format is accepted and gives us a real "
              "leaderboard number to compare against. Getting a format error on the last night "
              "is the failure mode we are designing around.")

    head(doc, "6. One trap we are planning for")
    para(doc, "The training data only covers the US and India. The test set adds France, which "
              "we have never seen. If anything in our code quietly learns US or Indian address "
              "habits, we lose 15 percent of the final score. Our check is cheap: train on US "
              "records only, test on Indian records only. If that collapses, we have a problem "
              "worth fixing before day three.")

    head(doc, "7. What to pick up")
    para(doc, "Claim a phase in the group chat before you start so we do not build the same "
              "thing twice. Full detail for every phase is in the README on GitHub.")
    bullet(doc, "Steps 1 and 2 are the bottleneck. Two people on blocking is not wasted effort.")
    bullet(doc, "Log every run in experiments/experiments.md, including the ones that fail. "
                "If it is not logged we cannot compare it or write it up later.")
    bullet(doc, "Nothing gets uploaded unless it beats our best local score and the official "
                "validator passes. We only get five uploads a day.")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    r = note.add_run("Everything above is built from the provided training data only. No outside "
                     "databases, no lookup services, no geocoding APIs. That is a disqualification "
                     "rule, and it is worth all of us knowing it.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
